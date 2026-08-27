#!/usr/bin/env python3
"""
Fetch official Government of India energy statistics and emit a JSON feed
for the frontend: frontend/src/data/officialStats.json

Sources (official GoI publications, no third-party aggregators):
  1. MNRE Physical Progress        -> monthly state-wise RE installed capacity PDF
                                      https://mnre.gov.in/en/physical-progress/
  2. CEA Installed Capacity Report -> monthly XLSX, all-India capacity mix by mode
                                      https://cea.nic.in/installed-capacity-report?lang=en
  3. CEA Renewable Generation      -> monthly state-wise RE generation XLSX
                                      https://cea.nic.in/renewable-generation-report/?lang=en

Each source is independently optional: if one fails to download or parse, the
script still emits valid JSON with the remaining sections and records a warning
in meta.warnings (it never fabricates data).

Usage:
    python scripts/fetch_official_data.py                     # discover latest + download
    python scripts/fetch_official_data.py --refresh           # force re-download
    python scripts/fetch_official_data.py --offline           # use cached files only
    python scripts/fetch_official_data.py --mnre-pdf FILE \
        --cea-ic-xlsx FILE --cea-re-xlsx FILE                 # explicit local files
"""
import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import requests

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = PROJECT_ROOT / "frontend" / "src" / "data" / "officialStats.json"
DEFAULT_CACHE = PROJECT_ROOT / "data" / "official"

MNRE_PAGE = "https://mnre.gov.in/en/physical-progress/"
CEA_IC_PAGE = "https://cea.nic.in/installed-capacity-report?lang=en"
CEA_RE_PAGE = "https://cea.nic.in/renewable-generation-report/?lang=en"

HTTP_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) SolarSite-India research fetcher"
    )
}


def http_get(url, timeout=60):
    resp = requests.get(url, headers=HTTP_HEADERS, timeout=timeout)
    resp.raise_for_status()
    return resp


def download(url, cache_dir, refresh=False, name=None):
    cache_dir.mkdir(parents=True, exist_ok=True)
    fname = name or Path(urlparse(url).path).name or "download.bin"
    dest = cache_dir / fname
    if dest.exists() and not refresh:
        print(f"  [cache] {fname}")
        return dest
    print(f"  [get ] {url[:100]}")
    dest.write_bytes(http_get(url).content)
    print(f"  [saved] {dest.name} ({dest.stat().st_size:,} bytes)")
    return dest


def clean_number(value):
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("\n", "")
    if text in {"", "-", "--", "NA", "N/A", "Nil", "NIL"}:
        return None
    try:
        return round(float(text), 4)
    except ValueError:
        return None


def normalize_state(name):
    if not name:
        return None
    text = str(name).strip()
    if "/" in text and re.search(r"[\u0900-\u097F]", text.split("/")[0]):
        text = text.split("/")[-1]
    text = re.sub(r"\s+", " ", text).strip(" *#")
    fixes = {
        "Jammu & Kashmir": "Jammu and Kashmir",
        "Jammu & Kashmir (including Ladakh)": "Jammu and Kashmir",
        "Jammu Kashmir": "Jammu and Kashmir",
        "Ladakh": "Ladakh",
        "Dadra & Nagar Haveli and Daman & Diu": "Dadra and Nagar Haveli and Daman and Diu",
        "Daman & Diu": "Dadra and Nagar Haveli and Daman and Diu",
        "Andaman & Nicobar Islands": "Andaman and Nicobar Islands",
        "Andaman & Nicobar": "Andaman and Nicobar Islands",
        "Andaman & Nicobar I": "Andaman and Nicobar Islands",
        "Pudduchery": "Puducherry",
    }
    return fixes.get(text, text)


def discover_mnre_pdf():
    page = http_get(MNRE_PAGE).text
    candidates = []
    for match in re.finditer(r"<a\b[^>]*>", page):
        tag = match.group(0)
        href = re.search(r'href="(https://cdnbbsr[^"]+\.pdf)"', tag)
        if not href:
            continue
        label_match = re.search(r'aria-label="([^"]*)"', tag) or re.search(r">\s*([^<]*)", tag)
        label = label_match.group(1) if label_match else ""
        if "state wise re installed capacity" in label.lower():
            candidates.append((href.group(1), label))
    if not candidates:
        raise RuntimeError("could not locate state-wise capacity PDF link on MNRE page")
    url, label = candidates[0]
    date_match = re.search(r"as on\s+(\d{2}\.\d{2}\.\d{4})", label)
    as_on = date_match.group(1) if date_match else None
    return url, {"asOn": as_on}


def parse_mnre_pdf(path):
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError(f"pdfplumber is required to parse the MNRE PDF ({exc})")

    columns = [
        "solarGroundMountedMW", "solarRooftopMW", "solarHybridMW", "solarOffgridMW",
        "solarTotalMW", "windMW", "biomassBagasseMW", "biomassCogenMW",
        "wasteToEnergyMW", "wasteToEnergyOffgridMW", "bioPowerTotalMW",
        "smallHydroMW", "largeHydroMW", "totalREMW",
    ]
    states, all_india = [], {}
    with pdfplumber.open(path) as pdf:
        tables = pdf.pages[0].extract_tables()
        if not tables:
            raise RuntimeError("no tables found in MNRE PDF")
        for row in tables[0]:
            cells = [re.sub(r"\s+", " ", c).strip() if c else "" for c in row]
            label = cells[2] if len(cells) > 2 else ""
            numbers = [clean_number(c) for c in cells[3:17]]
            if len(numbers) < len(columns):
                continue
            if label.startswith("Total"):
                all_india = dict(zip(columns, numbers))
                break
            if label.startswith("S. No") or label in {"", "STATES / UTs"} or "(MW)" in label:
                continue
            if not label or cells[0] == "":
                continue
            if not re.match(r"^\d+$", cells[0]):
                continue
            states.append({"state": normalize_state(label), **dict(zip(columns, numbers))})
    if not states:
        raise RuntimeError("MNRE PDF parsed but yielded no state rows")
    return states, all_india


def parse_cea_ic_xlsx(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(f"openpyxl is required to parse the CEA capacity XLSX ({exc})")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    as_on = None
    for row in rows[:5]:
        for cell in row:
            if cell and "As on" in str(cell):
                m = re.search(r"As on\s+(\d{2}\.\d{2}\.\d{4})", str(cell))
                if m:
                    as_on = m.group(1)

    totals = None
    for i, row in enumerate(rows):
        region = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        sector = str(row[2]).strip().lower() if len(row) > 2 and row[2] else ""
        if region.upper() == "ALL INDIA":
            for j in range(i, min(i + 6, len(rows))):
                if len(rows[j]) > 12 and str(rows[j][2]).strip().lower() in {"total", "sub total"}:
                    r = rows[j]
                    totals = {
                        "coalMW": clean_number(r[3]),
                        "ligniteMW": clean_number(r[4]),
                        "gasMW": clean_number(r[5]),
                        "dieselMW": clean_number(r[6]),
                        "thermalTotalMW": clean_number(r[7]),
                        "nuclearMW": clean_number(r[8]),
                        "hydroMW": clean_number(r[9]),
                        "renewableMW": clean_number(r[10]),
                        "renewableInclHydroMW": clean_number(r[11]),
                        "grandTotalMW": clean_number(r[12]),
                    }
                    break
            break
    if totals is None:
        raise RuntimeError("could not find ALL INDIA total row in CEA capacity XLSX")
    return totals, as_on


def parse_cea_re_gen(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(f"openpyxl is required to parse the CEA generation XLSX ({exc})")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = next((s for s in wb.sheetnames if s.lower().startswith("state wise")), None)
    if sheet_name is None:
        wb.close()
        raise RuntimeError("no 'State Wise' sheet found in CEA generation XLSX")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))

    header_idx, periods = None, []
    for i, row in enumerate(rows[:20]):
        joined = [str(c) for c in row if c]
        if any("Name of State/UT" in c for c in joined):
            header_idx = i
            for cell in row[2:6]:
                if cell:
                    periods.append(re.sub(r"\s+", " ", str(cell)).strip())
            break
    if header_idx is None:
        wb.close()
        raise RuntimeError("header row not found in CEA generation sheet")

    skip_markers = ("क्षेत्र", "Sub To", "Region")
    all_india_row, states = None, []
    for row in rows[header_idx + 1:]:
        name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if not name or name.startswith(("Details", "Note", "1.", "2.", "*")):
            continue
        values = [clean_number(row[k]) if len(row) > k else None for k in range(2, 6)]
        if "सम्पूर्ण भारत" in name:
            all_india_row = values
            continue
        if any(m in name for m in skip_markers):
            continue
        english = normalize_state(name)
        if not english or values[0] is None:
            continue
        states.append({
            "state": english,
            "genCurrentMonthMU": values[0],
            "genSameMonthPrevYearMU": values[1],
            "genFytdCurrentMU": values[2],
            "genFytdPrevYearMU": values[3],
        })
    wb.close()
    if not states:
        raise RuntimeError("CEA generation sheet parsed but yielded no state rows")
    return states, (all_india_row or []), periods


def merge_states(capacity_rows, generation_rows):
    gen_by_state = {g["state"]: g for g in generation_rows}
    merged = []
    for cap in capacity_rows:
        entry = dict(cap)
        gen = gen_by_state.pop(entry["state"], None)
        if gen:
            entry.update({k: v for k, v in gen.items() if k != "state"})
        merged.append(entry)
    for leftover in gen_by_state.values():
        merged.append({
            "state": leftover["state"],
            "genCurrentMonthMU": leftover.get("genCurrentMonthMU"),
            "genSameMonthPrevYearMU": leftover.get("genSameMonthPrevYearMU"),
            "genFytdCurrentMU": leftover.get("genFytdCurrentMU"),
            "genFytdPrevYearMU": leftover.get("genFytdPrevYearMU"),
        })
    priority = {"andhra pradesh", "telangana"}
    merged.sort(key=lambda s: (s["state"] not in priority, -(s.get("totalREMW") or s.get("genCurrentMonthMU") or 0)))
    return merged


def main():
    parser = argparse.ArgumentParser(description="Fetch official GoI energy stats -> officialStats.json")
    parser.add_argument("--output", type=str, default=str(DEFAULT_OUTPUT))
    parser.add_argument("--cache-dir", type=str, default=str(DEFAULT_CACHE))
    parser.add_argument("--mnre-pdf", type=str, default=None, help="local MNRE state-wise PDF")
    parser.add_argument("--cea-ic-xlsx", type=str, default=None, help="local CEA installed capacity XLSX")
    parser.add_argument("--cea-re-xlsx", type=str, default=None, help="local CEA RE generation XLSX")
    parser.add_argument("--offline", action="store_true", help="use only cached files")
    parser.add_argument("--refresh", action="store_true", help="force re-download")
    args = parser.parse_args()

    cache_dir = Path(args.cache_dir)
    warnings = []
    sources = {}

    def resolve(local_arg, discover_fn, cache_glob, label, download_name=None):
        if local_arg:
            path = Path(local_arg)
            if not path.exists():
                raise FileNotFoundError(path)
            return path, None
        cached = sorted(cache_dir.glob(cache_glob))
        if cached and not args.refresh:
            return cached[-1], None
        if args.offline:
            raise RuntimeError(f"--offline set and no cached file matching {cache_glob}")
        url, extra = discover_fn()
        return download(url, cache_dir, refresh=args.refresh, name=download_name), {"url": url, **(extra or {})}

    capacity_rows, all_india_re = [], {}
    try:
        print("[1/3] MNRE state-wise RE installed capacity")
        if args.mnre_pdf:
            path = Path(args.mnre_pdf)
            if not path.exists():
                raise FileNotFoundError(path)
            info = {}
        else:
            cached = sorted(cache_dir.glob("mnre_statewise_*.pdf"))
            if cached and not args.refresh:
                path = cached[-1]
                m = re.search(r"(\d{2}\.\d{2}\.\d{4})", path.name)
                info = {"asOn": m.group(1) if m else None, "url": None}
            else:
                if args.offline:
                    raise RuntimeError("--offline set and no cached MNRE PDF")
                url, extra = discover_mnre_pdf()
                name = f"mnre_statewise_{extra.get('asOn') or 'latest'}.pdf"
                path = download(url, cache_dir, refresh=args.refresh, name=name)
                info = {"url": url, **extra}
        states, totals = parse_mnre_pdf(path)
        capacity_rows, all_india_re = states, totals
        sources["stateCapacity"] = {
            "publisher": "MNRE", "file": path.name,
            "asOn": (info or {}).get("asOn"), "url": (info or {}).get("url"),
        }
        print(f"  parsed {len(states)} states, national solar "
              f"{(totals.get('solarTotalMW') or 0)/1000:.2f} GW")
    except Exception as exc:
        warnings.append(f"stateCapacity unavailable: {exc}")
        print(f"  WARNING: {exc}")

    capacity_mix, ic_as_on = {}, None
    try:
        print("[2/3] CEA installed capacity mix")

        def discover_ic():
            page = http_get(CEA_IC_PAGE).text
            m = re.search(r'href="(https://cea\.nic\.in/wp-content/uploads/installed/[^"]+\.xlsx)"', page)
            if not m:
                raise RuntimeError("no installed-capacity XLSX link on CEA page")
            return m.group(1), {}

        path, info = resolve(args.cea_ic_xlsx, discover_ic, "IC_*.xlsx", "cea-ic")
        capacity_mix, ic_as_on = parse_cea_ic_xlsx(path)
        sources["capacityMix"] = {"publisher": "CEA", "file": path.name,
                                  "asOn": ic_as_on, "url": (info or {}).get("url")}
        print(f"  grand total {capacity_mix.get('grandTotalMW'):,.0f} MW")
    except Exception as exc:
        warnings.append(f"capacityMix unavailable: {exc}")
        print(f"  WARNING: {exc}")

    generation_rows, gen_all_india, periods = [], [], []
    try:
        print("[3/3] CEA state-wise RE generation")

        def discover_re_gen():
            page = http_get(CEA_RE_PAGE).text
            links = re.findall(
                r'href="(https://cea\.nic\.in/wp-content/uploads/resd/[^"]+/RE_Monthly_Generation_report_[^"]+\.xlsx)"',
                page,
            )
            if not links:
                raise RuntimeError("no RE generation XLSX link on CEA page")
            return links[0], {}

        path, info = resolve(args.cea_re_xlsx, discover_re_gen, "RE_Monthly_Generation_*.xlsx", "cea-re")
        generation_rows, gen_all_india, periods = parse_cea_re_gen(path)
        sources["reGeneration"] = {
            "publisher": "CEA", "file": path.name, "periods": periods,
            "url": (info or {}).get("url"),
        }
        print(f"  parsed {len(generation_rows)} states/UTs")
    except Exception as exc:
        warnings.append(f"reGeneration unavailable: {exc}")
        print(f"  WARNING: {exc}")

    states_merged = merge_states(capacity_rows, generation_rows)
    priority = {s["state"]: s for s in states_merged if s["state"] in {"Andhra Pradesh", "Telangana"}}

    feed = {
        "meta": {
            "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "generator": "scripts/fetch_official_data.py",
            "sources": sources,
            "coverage": {
                "statesWithCapacity": sum(1 for s in states_merged if s.get("totalREMW") is not None),
                "statesWithGeneration": sum(1 for s in states_merged if s.get("genCurrentMonthMU") is not None),
            },
            "notes": [
                "All figures from official Government of India publications (MNRE / CEA).",
                "Capacities in MW; generation in MU (million units = GWh).",
                "Empty numeric fields mean the source reported no value for that cell.",
            ],
            "warnings": warnings,
        },
        "allIndia": {
            "reCapacity": all_india_re or None,
            "generation": {
                "currentMonthMU": gen_all_india[0] if len(gen_all_india) > 0 else None,
                "sameMonthPrevYearMU": gen_all_india[1] if len(gen_all_india) > 1 else None,
                "fytdCurrentMU": gen_all_india[2] if len(gen_all_india) > 2 else None,
                "fytdPrevYearMU": gen_all_india[3] if len(gen_all_india) > 3 else None,
            } if gen_all_india else None,
            "installedCapacityMix": capacity_mix or None,
        },
        "priorityStates": priority,
        "states": states_merged,
    }

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(feed, indent=2, ensure_ascii=False))
    print(f"\nWrote {out_path} ({out_path.stat().st_size:,} bytes)")
    if warnings:
        print(f"{len(warnings)} warning(s): see meta.warnings")
    for state, row in priority.items():
        print(f"  {state}: solar {row.get('solarTotalMW', 0):,.0f} MW | "
              f"total RE {row.get('totalREMW', 0):,.0f} MW | "
              f"gen {row.get('genCurrentMonthMU') or 0:,.0f} MU")
    return 0


if __name__ == "__main__":
    sys.exit(main())
